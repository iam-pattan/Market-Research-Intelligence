"""Build the project walkthrough as a Word document (docs/How_the_Banyan_Screen_Was_Built.docx).

Same nine-stage narrative as reports/project_walkthrough.html; charts are
rendered with matplotlib and embedded as images. Figures are the same ones
the HTML page carries (data/ + docs/TOOLING_LOG.md + docs/HANDOFF.md).

Usage: python analysis/build_walkthrough_docx.py
"""
from __future__ import annotations

import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "How_the_Banyan_Screen_Was_Built.docx"
TMP = Path(tempfile.mkdtemp(prefix="walkthrough_"))

ACCENT, LIGHT, NEUTRAL = "#0f6e75", "#9ccdd2", "#c9c5b9"
GOOD, WARN, CRIT = "#0ca30c", "#fab219", "#d03b3b"
S1, S2 = "#2a78d6", "#eb6834"
INK, INK2 = "#1c1b18", "#5b5852"

plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 10, "axes.edgecolor": "#c9c5b9",
    "axes.labelcolor": INK2, "xtick.color": INK2, "ytick.color": INK, "text.color": INK,
    "axes.spines.top": False, "axes.spines.right": False, "axes.spines.left": False,
    "figure.dpi": 200, "savefig.dpi": 200, "savefig.bbox": "tight", "savefig.facecolor": "white",
})


# ----------------------------------------------------------------- charts ---
def hbar(name, labels, values, colors=None, xlabel="", fmt="{:g}", height=None, note=None):
    n = len(labels)
    fig, ax = plt.subplots(figsize=(7.2, height or (0.42 * n + 0.9)))
    colors = colors or [ACCENT] * n
    y = list(range(n))[::-1]
    ax.barh(y, values, color=colors, height=0.62)
    ax.set_yticks(y)
    ax.set_yticklabels(labels)
    ax.tick_params(axis="y", length=0)
    ax.set_xlim(0, max(values) * 1.18)
    ax.xaxis.grid(True, color="#e6e3db", linewidth=0.6)
    ax.set_axisbelow(True)
    for yi, v in zip(y, values):
        ax.text(v + max(values) * 0.015, yi, fmt.format(v), va="center", fontsize=9, color=INK2)
    if xlabel:
        ax.set_xlabel(xlabel)
    if note:
        fig.text(0.01, -0.02, note, fontsize=8, color=INK2)
    p = TMP / f"{name}.png"
    fig.savefig(p)
    plt.close(fig)
    return p


def qc_chart():
    cats = ["Review score", "GLM 5.2", "Kimi 2.7"]
    r1, r2 = [65.9, 74.6, 57.2], [66.9, 81.0, 52.8]
    fig, ax = plt.subplots(figsize=(7.2, 2.6))
    y = [2, 1, 0]
    h = 0.34
    ax.barh([v + h / 2 + 0.03 for v in y], r1, height=h, color=S1, label="Round 1")
    ax.barh([v - h / 2 - 0.03 for v in y], r2, height=h, color=S2, label="Round 2")
    for yi, a, b in zip(y, r1, r2):
        ax.text(a + 1, yi + h / 2 + 0.03, f"{a}", va="center", fontsize=9, color=INK2)
        ax.text(b + 1, yi - h / 2 - 0.03, f"{b}", va="center", fontsize=9, color=INK2)
    ax.set_yticks(y)
    ax.set_yticklabels(cats)
    ax.tick_params(axis="y", length=0)
    ax.set_xlim(0, 100)
    ax.xaxis.grid(True, color="#e6e3db", linewidth=0.6)
    ax.set_axisbelow(True)
    ax.set_xlabel("score / 100")
    ax.legend(frameon=False, loc="lower right", fontsize=9)
    p = TMP / "qc.png"
    fig.savefig(p)
    plt.close(fig)
    return p


def flow_chart():
    steps = [("Workflow script", "analysis/workflows/*.js\npipeline / parallel, resumable"),
             ("N research agents", "Claude Code subagents\nrunning in parallel"),
             ("codex --search", "Codex CLI via Bash, web search on\nChatGPT-authed, no API key"),
             ("Cited findings", "raw research text saved\n(the audit trail)"),
             ("Synthesis agent", "JSON-schema-forced\nstructured output"),
             ("data/*.json", "plain files, no hidden state\nre-runnable from disk")]
    fig, ax = plt.subplots(figsize=(7.6, 3.6))
    ax.axis("off")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    cols = 3
    w, hgt, gap = 0.29, 0.36, 0.055
    xs = [0.02 + c * (w + gap) for c in range(cols)]
    ys = [0.58, 0.06]
    centers = []
    for i, (t, sub) in enumerate(steps):
        c, r = i % cols, i // cols
        x, y = xs[c], ys[r]
        hi = i == 2
        ax.add_patch(plt.Rectangle((x, y), w, hgt, facecolor="#dcecee" if hi else "#f1efe8",
                                   edgecolor=ACCENT if hi else "#c9c5b9", linewidth=1.4 if hi else 0.9,
                                   linestyle="--" if i == 5 else "-"))
        ax.text(x + 0.018, y + hgt - 0.05, f"{i + 1}", fontsize=8, color=INK2, va="top")
        ax.text(x + 0.05, y + hgt - 0.06, t, fontsize=10, weight="bold", va="top")
        ax.text(x + 0.05, y + hgt - 0.17, sub, fontsize=7, color=INK2, va="top", linespacing=1.3)
        centers.append((x + w / 2, y + hgt / 2, x, y))
    arr = dict(arrowstyle="->", color=INK2, lw=0.9, shrinkA=0, shrinkB=0)
    for i in range(len(steps) - 1):
        cx, cy, x, y = centers[i]
        nx, ny, x2, y2 = centers[i + 1]
        if i % cols < cols - 1:  # same row: right edge -> left edge
            ax.annotate("", xy=(x2, cy), xytext=(x + w, cy), arrowprops=arr)
        else:  # row break: down from box 3, left along the gutter, down into box 4
            ymid = (y + y2 + hgt) / 2
            ax.plot([cx, cx, nx], [y, ymid, ymid], color=INK2, lw=0.9)
            ax.annotate("", xy=(nx, y2 + hgt), xytext=(nx, ymid), arrowprops=arr)
    p = TMP / "flow.png"
    fig.savefig(p)
    plt.close(fig)
    return p


def funnel_chart():
    labels = ["Records discovered", "After dedup (best-sourced wins)", "Curated top-N (best-of-both)",
              "Deep-researched (Banyan top-25)", "Outreach recommended (SEND)"]
    vals = [333, 318, 250, 25, 18]
    colors = [ACCENT, ACCENT, LIGHT, LIGHT, GOOD]
    return hbar("funnel", labels, vals, colors, xlabel="companies")


# ------------------------------------------------------------------ docx ---
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


class Doc:
    def __init__(self):
        self.d = Document()
        st = self.d.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        for lvl, size in ((1, 20), (2, 14), (3, 12)):
            h = self.d.styles[f"Heading {lvl}"]
            h.font.name = "Calibri"
            h.font.size = Pt(size)
            h.font.color.rgb = RGBColor(0x0F, 0x6E, 0x75) if lvl == 1 else RGBColor(0x1C, 0x1B, 0x18)
        for s in self.d.sections:
            s.left_margin = s.right_margin = Inches(1)
            s.top_margin = s.bottom_margin = Inches(0.9)

    def h(self, text, lvl=1):
        return self.d.add_heading(text, lvl)

    def p(self, text="", italic=False, size=None, color=None, bold=False, after=6):
        para = self.d.add_paragraph()
        run = para.add_run(text)
        run.italic, run.bold = italic, bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        para.paragraph_format.space_after = Pt(after)
        return para

    def rich(self, parts, after=6):
        """parts: list of (text, {'b':bool,'i':bool,'code':bool})"""
        para = self.d.add_paragraph()
        for text, fmt in parts:
            r = para.add_run(text)
            r.bold = fmt.get("b", False)
            r.italic = fmt.get("i", False)
            if fmt.get("code"):
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)
        para.paragraph_format.space_after = Pt(after)
        return para

    def bullets(self, items):
        for it in items:
            para = self.d.add_paragraph(style="List Bullet")
            if isinstance(it, tuple):
                r = para.add_run(it[0])
                r.bold = True
                para.add_run(it[1])
            else:
                para.add_run(it)
            para.paragraph_format.space_after = Pt(3)

    def fig(self, path, caption, width=6.3):
        self.d.add_picture(str(path), width=Inches(width))
        self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = self.p(caption, italic=True, size=9, color=INK2, after=12)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def table(self, header, rows, widths=None, mono_cols=(), status_col=None):
        t = self.d.add_table(rows=1, cols=len(header))
        t.style = "Light Grid Accent 1"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, htxt in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            r = cell.paragraphs[0].add_run(htxt)
            r.bold = True
            r.font.size = Pt(9.5)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(str(val))
                r.font.size = Pt(9.5)
                if i in mono_cols:
                    r.font.name = "Consolas"
                    r.font.size = Pt(8.5)
                if status_col is not None and i == status_col:
                    v = str(val).lower()
                    fill = {"used": "E3F4E3", "fired": "E3F4E3", "silent": "E3F4E3", "built here": "E3F4E3",
                            "load-bearing": "E3F4E3", "this session": "FFF3D6", "pattern only": "DCECEE"}.get(v)
                    if v.startswith("failed"):
                        fill = "F9E0E0"
                    if fill:
                        shade(cells[i], fill)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def stage(self, n, label, title, why):
        self.d.add_page_break()
        self.p(f"STAGE {n} · {label.upper()}", size=9, bold=True, color=ACCENT, after=0)
        self.h(title, 1)
        self.p(why, italic=True, color=INK2, after=10)

    def save(self, path):
        self.d.save(path)


def main() -> None:
    D = Doc()

    # ---- title
    D.p("PROJECT WALKTHROUGH · BANYAN M&A SCREENING", size=9, bold=True, color=ACCENT, after=2)
    t = D.d.add_heading("How the Banyan screen was built, and why it looks the way it does", 0)
    t.runs[0].font.size = Pt(24)
    t.runs[0].font.color.rgb = RGBColor(0x1C, 0x1B, 0x18)
    D.p("A guided tour of the thinking behind the pipeline: which sources were trusted, how the data was "
        "pulled without a single paid database, which agents, skills and hooks did the work, and how every "
        "number was checked before it reached a deliverable.", size=12, color=INK2, after=10)
    D.table(["Companies discovered", "Curated & scored", "Deep dossiers", "Triangulated universe",
             "Multi-agent workflows", "Unit tests", "Paid data sources"],
            [[333, 250, 25, 602, 4, 53, 0]])
    D.h("Four constraints that shaped every decision", 2)
    D.bullets([
        ("Public web + registries only. ", "No PitchBook, Crunchbase or Grata. Every fact traces to a company site, a filed government record, or a cited search result."),
        ("Uncertainty must be visible. ", "Private-company financials are mostly unknowable from outside. The design carries a confidence on every signal and discounts scores instead of guessing."),
        ("Discovery-first. ", "There was no seed list. The system had to find the universe before it could rank it."),
        ("Four audiences, four entitlements. ", "Sales, Finance, Data Science and Leadership each see a different slice. Access control is part of the output, not an afterthought."),
    ])

    # ---- 1 frame
    D.stage(1, "Frame", "Turn the thesis into a rubric before touching any data",
            "Banyan's buy criteria are public and specific. Writing them down as weighted, gated criteria first meant "
            "every later step had a fixed target to score against — and the LLM never got to decide what \"good\" means.")
    D.rich([("The work opened with the ", {}), ("superpowers:brainstorming", {"code": True}),
            (" skill (scope the pipeline and the IAM model with the user, one question at a time) and ", {}),
            ("superpowers:writing-plans", {"code": True}),
            (" (a task-by-task plan for the scoring core). The rubric became ", {}), ("config.yaml", {"code": True}),
            (": seven criteria with weights derived from banyansoftware.com's stated FAQ, two hard gates, a confidence floor "
             "and tier thresholds. A second, contrasting rubric (", {}), ("config.growth.yaml", {"code": True}),
            (") exists purely to catch companies that look impressive but are the wrong shape — VC-backed, scaling, unprofitable.", {})])
    D.fig(hbar("rubric", ["Recurring revenue", "Profitability", "Niche leadership", "Revenue band ($2M–$100M)",
                          "Customer retention", "Ownership fit (founder/family)", "Team stability"],
               [0.25, 0.20, 0.20, 0.10, 0.10, 0.10, 0.05], fmt="{:.2f}", xlabel="weight"),
          "Figure 1 — The Banyan rubric as weights. Gates apply before the tier: loss evidence → rejected; revenue under $2M → capped at tier C.")

    # ---- 2 sources
    D.stage(2, "Sources", "Three independent ways of finding companies",
            "One discovery method has blind spots. Three methods that do not share a source can be cross-checked against "
            "each other — that is where \"founder-owned\" claims get verified rather than believed.")
    D.fig(hbar("passes", ["Pass 1 · Vertical discovery", "Pass 2 · Acquirer adjacency", "Pass 3 · OSINT & registries"],
               [318, 285, 152], xlabel="candidates surfaced (before de-duplication)"),
          "Figure 2 — Candidates per discovery pass. The union across all three is 602 companies.")
    D.bullets([
        ("Pass 1 ", "searched the public web vertical by vertical — 12 batches: construction, healthcare practice management, K-12 systems, manufacturing ERP and so on."),
        ("Pass 2 ", "asked a different question: what do Constellation, Valsoft and Everfield already own, and who are the peers of those companies? Ten angles."),
        ("Pass 3 ", "avoided search engines altogether and went to primary records — registries, filed accounts, GitHub, hiring pages, product directories."),
    ])
    D.fig(hbar("osint", ["GitHub tech footprint", "Product directory leaders", "ANZ / Canada registries", "Careers & hiring signals",
                         "US registries / SEC EDGAR", "EU registries (filed accounts)", "UK / Ireland Companies House"],
               [31, 31, 24, 22, 20, 19, 5], colors=[ACCENT] * 5 + [GOOD] * 2, xlabel="candidates",
               note="Green = filed financial statements available — the only place a private company's real revenue and profit exist."),
          "Figure 3 — Inside pass 3: the seven OSINT angles. Pass 3 also ran 34 ownership verifications against registers, "
          "which caught 11 companies described as \"founder-owned\" on their own websites that are actually PE- or VC-controlled.")
    D.h("What was deliberately not used", 2)
    D.table(["Tool", "Role", "Status"], [
        ["PitchBook · Crunchbase · Grata", "Paid private-company databases", "Excluded by design"],
        ["perplexity-research", "Attempted as a discovery engine", "Failed — no API key"],
        ["firecrawl-scrape", "Attempted as a scraping engine", "Failed — CLI not authenticated"],
        ["WebSearch (built-in)", "Blocked in the sandboxed environment", "Failed — unavailable"],
    ], widths=[2.0, 2.9, 1.6], mono_cols=(0,), status_col=2)

    # ---- 3 mechanism
    D.stage(3, "Mechanism", "How the data was actually pulled",
            "The engine is the OpenAI Codex CLI with web search turned on, driven by Claude Code's multi-agent Workflow tool. "
            "Every research call returns cited findings that a second agent turns into structured JSON under a schema — so the raw "
            "research is kept as an audit trail and the structured record can be checked against it later.")
    D.fig(flow_chart(), "Figure 4 — One research workflow, end to end. The same shape ran for every research stage with different prompts and fan-out sizes.")
    D.rich([("The ", {}), ("ai-hub-web-search:codex-websearch", {"code": True}),
            (" skill supplied the pattern (dispatch search from a subagent, sanitise the query, cite sources). Its shipped script was "
             "bypassed — it hard-coded invalid model ids and a blocked /tmp path — so the CLI was called directly with corrected flags, "
             "in the foreground with the sandbox disabled. ", {}), ("workflow-authoring", {"code": True}),
            (" was loaded before every workflow was written.", {})])
    D.h("The workflows", 2)
    D.table(["Workflow", "Fan-out", "Question it answered", "Output"], [
        ["vertical discovery (12 batches)", "12 batches", "Which vertical software companies exist in each niche, with per-criterion signals and confidence?", "data/records/*.json (333)"],
        ["independent_discovery_wf.js", "10 angles", "Who do the permanent-hold consolidators already own, and who are their peers?", "independent_results.json (285)"],
        ["osint_discovery_wf.js", "7 angles + verify", "What do registries, filed accounts, GitHub and hiring data say — and is \"founder-owned\" true?", "osint_results.json (152 + 34 checks)"],
        ["market_synthesis_wf.js", "29 agents", "What does each of the 14 macro-segments look like as a hunting ground?", "market_synthesis.json"],
        ["banyan25_dossiers_wf.js", "50 agents", "For each top-25 target: business model, ownership, moat, risks, sources — and a thesis-appropriate outreach email.", "raw_dossiers_banyan.json"],
    ], widths=[1.7, 0.9, 2.6, 1.3], mono_cols=(0, 3))

    # ---- 4 score
    D.stage(4, "Score", "A deterministic core the LLM cannot touch",
            "The research agents propose signals with confidence; a hand-written, unit-tested Python package (banyan_screen/) turns them "
            "into scores and tiers. No model call happens inside the scoring — so the same record always produces the same tier, and every "
            "result carries a rationale you can audit.")
    D.fig(funnel_chart(), "Figure 5 — From discovery to the ranked list. Each bar is a step in banyan_screen.run.")
    D.fig(hbar("tiers", ["Tier B", "Tier C", "Rejected (loss evidence)", "Insufficient data"], [23, 202, 24, 1],
               colors=[GOOD, NEUTRAL, CRIT, WARN], xlabel="companies (of 250)"),
          "Figure 6 — Where the 250 landed under the Banyan rubric. Tier C is the expected home of a genuine private target whose books are not public — it means \"verify\", not \"reject\".")
    D.rich([("adjusted = composite × (0.5 + 0.5 × confidence)", {"code": True}),
            (" — a company nobody can verify loses at most half its composite score, never all of it.", {})])

    # ---- 5 triangulate
    D.stage(5, "Triangulate", "Cross-match the three passes",
            "Matching by domain and normalised name, build_crossmatch.py sorts the 602-company union into what is corroborated, "
            "what is new, and where the ownership story changes.")
    D.fig(hbar("xmatch", ["Union of all passes", "Found by pass 1 only", "Net-new from passes 2–3",
                          "…of which strong (founder-owned or filed accounts)", "Corroborated by ≥ 2 passes", "False \"founder-owned\" caught"],
               [602, 240, 284, 176, 90, 11], colors=[LIGHT, LIGHT, ACCENT, ACCENT, GOOD, CRIT], xlabel="companies"),
          "Figure 7 — Outcome of the cross-match (data/research/crossmatch.json).")
    D.p("The 176 strong net-new names (Bromcom, CDL Group, Open Dental, Planning Center, GIRO, Maptek and the EU registry names among "
        "them) are the highest-value unfinished work: they have not yet been scored through the rubric.")

    # ---- 6 gate
    D.stage(6, "Gate", "Decide who sees what, in code",
            "The IAM model was designed conceptually first — 5 classification tiers, 14 roles, red-teamed from three angles — and then "
            "written as a real filter, banyan_screen/policy.py, that every artifact passes through. Deny by default: an unknown role or an "
            "unrecognised field gets nothing.")
    D.table(["Role", "Receives", "Never receives"], [
        ["Data Science", "Aggregates only", "Per-record output (refused outright), contacts, financials, deal intent"],
        ["Finance", "Financials, identity, signals", "Contacts, deal intent, outreach copy"],
        ["Sales / BD", "Identity, signals, contacts (masked), the outreach pitch", "Financials, deal intent"],
        ["Leadership", "Deal intent; contacts and financials masked", "Outreach copy, raw financials"],
        ["Governance / platform", "Nothing above Public", "Business content of any tier"],
    ], widths=[1.4, 2.6, 2.5])
    D.p("The companion page Layered Data Access Architecture shows the live policy per role and explains what a Claude user needs to "
        "sit behind this wall: an SSO identity and an enforcement point in front of the data — nothing inside the model.")

    # ---- 7 validate
    D.stage(7, "Validate", "Check the work with a different model, then check the model against the sources",
            "Two independent checks close the loop. A second-model QC judge scores the session's changes; a grounding script confirms "
            "that what the dossiers say actually appears in the raw research text.")
    D.fig(qc_chart(), "Figure 8 — Second-model QC, two rounds (validation-loop:validate; GLM 5.2 primary, Kimi 2.7 tie-breaker). "
                      "Round 2 followed fixes to a cross-match bug, the ownership classifier and the runtime IAM gate. "
                      "The judges disagreed by 28 points in round 2, so the composite is low-confidence — the disagreement is reported, not averaged away.")
    D.table(["Check", "Result", "What it means"], [
        ["Grounding (analysis/validate_grounding.py)", "409 / 415 facts (98.6%)", "Dossier claims trace to the raw research; the 6 unmatched were extractor false-positives, not invention"],
        ["Unit tests", "53 passing", "Scoring core, policy layer and page builders"],
        ["Page re-derivation (analysis/qc_consolidated_pages.py)", "0 mismatches", "Every figure on the consolidated pages recomputed from data/ and policy.py"],
        ["Cross-pass name check", "1 discrepancy caught", "A founder's name (QT9) was wrong in the OSINT pass — never trust a single-pass ownership name"],
    ], widths=[2.3, 1.4, 2.8])

    # ---- 8 tooling
    D.stage(8, "Tooling", "The skills, agents, hooks and tools behind it",
            "Everything below is from docs/TOOLING_LOG.md and the consolidation session. Status is honest: some things were used, "
            "some supplied only a pattern, some failed, and some were available but never called.")
    D.h("Skills (invoked through Claude Code's Skill tool)", 2)
    D.table(["Skill", "Purpose", "Status"], [
        ["superpowers:brainstorming", "Scoped the pipeline and the IAM model before any code", "used"],
        ["superpowers:writing-plans", "Task-by-task plan for the scoring core", "used"],
        ["workflow-authoring", "Reference for every multi-agent Workflow script (~10 loads)", "used"],
        ["ai-hub-web-search:codex-websearch", "The research-dispatch pattern; its script was bypassed", "pattern only"],
        ["banyan-sales-pitch", "The project's own skill: thesis-branched, IAM-gated outreach", "built here"],
        ["validation-loop:validate", "Second-model QC, two rounds (plus one in the consolidation session)", "used"],
        ["handoff", "Produced docs/HANDOFF.md", "used"],
        ["artifact-design · dataviz", "Design and chart rules for the consolidated pages", "this session"],
    ], widths=[2.2, 3.3, 1.0], mono_cols=(0,), status_col=2)
    D.h("Agents and orchestration", 2)
    D.table(["Component", "Role", "Status"], [
        ["Workflow tool", "~10 runs: pipeline/parallel fan-out with cached resume", "used"],
        ["Agent (general-purpose)", "Subagents executing the codex search pattern", "used"],
        ["validation-loop:qc-reviewer", "Independent second-model review of the consolidated pages", "this session"],
        ["codex CLI (gpt-5.6-sol, --search)", "The discovery engine — an external CLI, not a Claude tool", "load-bearing"],
    ], widths=[2.2, 3.3, 1.0], mono_cols=(0,), status_col=2)
    D.h("Hooks (automation that fired without being asked)", 2)
    D.table(["Hook", "Role", "Status"], [
        ["Stop-hook goal", "Nudged for visible build progress when the session goal was unmet", "fired"],
        ["PostToolUse marker (validation-loop)", "Tracks changed files so /validate reviews only this session's work", "silent"],
        ["ai-qc-guardrails auto-review", "Workspace-wide automatic Stop-hook QC; manual /validate used instead", "available, unused"],
    ], widths=[2.2, 3.3, 1.0], mono_cols=(0,), status_col=2)
    D.h("Built-in tools and MCP", 2)
    D.table(["Tool", "Role", "Status"], [
        ["Bash", "Every codex call, pytest, file operations — sandbox off where network was needed", "used"],
        ["Read · Write · Edit", "All source, config and docs", "used"],
        ["AskUserQuestion", "Scope decisions at each fork: sources, depth, rubric choice", "used"],
        ["MCP servers (original build)", "None — zero mcp__ calls; all external access via the codex CLI", "none"],
        ["playwright · context-mode MCP", "Screenshots for layout review; large-output analysis kept out of context", "this session"],
        ["Artifact tool", "Publishing the hub, the access page and the walkthrough", "this session"],
    ], widths=[2.2, 3.3, 1.0], mono_cols=(0,), status_col=2)

    # ---- 9 lessons
    D.stage(9, "Lessons", "What the build taught us",
            "The decisions that mattered most were not technical — they were about which question to ask and how much to trust an answer.")
    D.bullets([
        ("The rubric decides the outreach list. ", "Ranked best-of-both, the top-25 produced 2 SEND recommendations — dominated by VC-backed AI infrastructure. Ranked by the Banyan rubric alone, the same pipeline produced 18. The scoring lens is the product."),
        ("Tier C is not a bad grade. ", "Genuine private targets land at C by construction: their financials cannot be verified from outside, so confidence is low and the score is discounted. Low verified density in a segment is a verification artifact, not absence of fit."),
        ("\"Privately held\" is not \"founder-owned\". ", "Eleven companies said one thing on their websites and another in the registers. Ownership must be checked against a second source before anyone sends an email."),
        ("Keep the raw research. ", "Because every dossier's source text was saved, a script could later prove 98.6% of dossier facts were grounded — and show the 6 flags were the extractor's fault, not invention."),
        ("Record the failures. ", "Perplexity and Firecrawl failed on authentication and were abandoned for codex. The pivot is in the tooling log so the next person does not repeat it."),
        ("Separate judgment from arithmetic. ", "The LLM proposes signals with confidence; deterministic code computes the tier. That split is what makes every score reproducible and every rationale auditable."),
    ])

    # ---- map
    D.h("Where everything lives", 1)
    D.table(["Path", "What it is"], [
        ["banyan_screen/", "The deterministic core — models, rubric engine, ingest, run CLI, and policy.py (IAM). 53 tests in tests/."],
        ["config.yaml · config.growth.yaml", "The two rubrics: weights, gates, thresholds."],
        ["data/records/", "333 discovered company records from the 12 vertical batches."],
        ["data/screen/", "The scored, curated top-250 (results.json, ranked.csv)."],
        ["data/research/", "Passes 2–3, cross-match, dossiers, market synthesis, grounding results, and the verbatim codex research audit trail."],
        ["analysis/", "Reporting scripts and the four multi-agent workflow scripts; templates/ for the consolidated pages."],
        ["reports/", "Self-contained HTML deliverables: the intelligence hub, the access architecture, the walkthrough, and the earlier dashboards."],
        ["skills/banyan-sales-pitch/", "The reusable outreach skill (rules, templates, playbook)."],
        ["docs/", "HLD, LLD, tooling log, session handoff, this document, and the original design spec and plan."],
    ], widths=[2.2, 4.3], mono_cols=(0,))
    D.p("Companion pages: Banyan Target Intelligence (the one-stop hub) and Layered Data Access Architecture. Figures in this document "
        "are taken from the repository's data files and docs as of 2026-09-14; Banyan's \"120+ acquired, 0 sold\" is self-reported "
        "positioning, not audited.", italic=True, size=9, color=INK2)

    D.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
